from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"
OUTPUT_FILE = OUTPUT_DIR / "Bao_ve_do_an_FinnCenter_Chuc_nang_va_Co_che_hoat_dong.docx"

NAVY = "15324A"
TEAL = "0D9488"
LIGHT_TEAL = "E8F7F5"
BLUE = "2563EB"
LIGHT_BLUE = "EAF2FF"
ORANGE = "F59E0B"
LIGHT_ORANGE = "FFF7E6"
LIGHT_GRAY = "F3F6F8"
MID_GRAY = "64748B"
WHITE = "FFFFFF"
BLACK = "17202A"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs.get(edge)
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_REPEAT_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_together = value


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở file trong Microsoft Word."
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2, placeholder, fld_char3])


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        r_pr.append(underline_node)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_paragraph(doc: Document, text: str = "", style: str | None = None, *, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc: Document, items: Iterable[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_callout(doc: Document, title: str, body: str, fill=LIGHT_BLUE, title_color=BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    set_cell_border(
        cell,
        left={"val": "single", "sz": "18", "color": title_color},
        top={"val": "nil"},
        right={"val": "nil"},
        bottom={"val": "nil"},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(title_color)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_key_value_table(doc: Document, rows: list[tuple[str, str]], widths=(1.7, 5.8)) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (key, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        for cell in cells:
            set_cell_margins(cell)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "3", "color": "D6DEE5"},
                left={"val": "single", "sz": "3", "color": "D6DEE5"},
                bottom={"val": "single", "sz": "3", "color": "D6DEE5"},
                right={"val": "single", "sz": "3", "color": "D6DEE5"},
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], LIGHT_GRAY if index % 2 == 0 else "E9EEF2")
        set_cell_shading(cells[1], WHITE if index % 2 == 0 else "FAFCFD")
        key_run = cells[0].paragraphs[0].add_run(key)
        key_run.bold = True
        key_run.font.color.rgb = RGBColor.from_string(NAVY)
        cells[1].paragraphs[0].add_run(value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_matrix_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    header_cells = table.rows[0].cells
    set_REPEAT_table_header(table.rows[0])
    for i, header in enumerate(headers):
        set_cell_shading(header_cells[i], NAVY)
        set_cell_margins(header_cells[i])
        p = header_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            header_cells[i].width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_shading(cells[i], WHITE if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_margins(cells[i])
            set_cell_border(
                cells[i],
                top={"val": "single", "sz": "3", "color": "D6DEE5"},
                left={"val": "single", "sz": "3", "color": "D6DEE5"},
                bottom={"val": "single", "sz": "3", "color": "D6DEE5"},
                right={"val": "single", "sz": "3", "color": "D6DEE5"},
            )
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].paragraphs[0].add_run(value)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_feature_section(doc: Document, index: int, feature: dict) -> None:
    heading = doc.add_heading(f"{index}. {feature['title']}", level=2)
    set_keep_with_next(heading)
    add_key_value_table(
        doc,
        [
            ("Mục tiêu", feature["goal"]),
            ("Đối tượng", feature["actors"]),
            ("Điểm vào", feature["entry"]),
        ],
    )

    for label, key, color in (
        ("Luồng sử dụng", "flow", TEAL),
        ("Cơ chế hoạt động", "mechanism", BLUE),
        ("Quy tắc và kiểm soát", "rules", ORANGE),
    ):
        p = doc.add_paragraph()
        set_keep_with_next(p)
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(color)
        add_bullets(doc, feature[key])

    p = doc.add_paragraph()
    set_keep_with_next(p)
    r = p.add_run("Bằng chứng trong mã nguồn")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    for source in feature["sources"]:
        q = doc.add_paragraph(style="Code Path")
        q.add_run(source)

    add_callout(
        doc,
        "Câu nói bảo vệ đề xuất",
        feature["defense"],
        fill=LIGHT_TEAL,
        title_color=TEAL,
    )


FEATURES = [
    {
        "title": "Khu vực công khai và marketplace khóa học",
        "goal": "Giúp khách truy cập hiểu sản phẩm, khám phá khóa học và giảng viên trước khi đăng nhập.",
        "actors": "Khách, học viên, giảng viên, quản trị viên.",
        "entry": "/, /courses, /courses/[id], /teachers, /teachers/[id], /top-students, /about",
        "flow": [
            "Trang chủ lấy khóa học, giảng viên và các số liệu thật để tạo điểm vào cho hệ thống.",
            "Người dùng tìm kiếm, lọc theo ngôn ngữ, trình độ, loại khóa học/kỹ năng và sắp xếp theo giá hoặc tên.",
            "Trang chi tiết hợp nhất thông tin khóa, giảng viên, module, lesson, test, số học viên và review.",
            "Nếu đã ghi danh thì chuyển sang học; nếu chưa ghi danh thì bắt đầu luồng mua/đăng ký.",
        ],
        "mechanism": [
            "Next.js App Router dùng Server Component để đọc dữ liệu và Client Component cho lọc, modal, biểu mẫu.",
            "Chỉ khóa ACTIVE được hiển thị công khai; chủ sở hữu và admin có quyền preview khóa chưa công khai.",
            "Dữ liệu ngôn ngữ và nhãn hiển thị được chuẩn hóa trong các thư viện language-display/public-teachers.",
            "Bảng xếp hạng tổng hợp test đạt, enrollment, lượt dùng AI và LearningActivity.",
        ],
        "rules": [
            "Không tiết lộ nội dung học cho người chưa có quyền.",
            "Đánh dấu khóa đã ghi danh để tránh mua trùng.",
            "Chỉ giáo viên đang hoạt động và khóa hợp lệ mới được đưa vào danh sách công khai.",
        ],
        "sources": ["app/page.tsx", "app/courses/page.tsx", "app/courses/[id]/page.tsx", "lib/public-teachers.ts"],
        "defense": "Marketplace không chỉ là danh sách tĩnh. Nó là lớp đọc dữ liệu đã được lọc theo trạng thái nghiệp vụ, đồng thời điều hướng đúng người dùng sang mua khóa, học tiếp hoặc preview theo vai trò.",
    },
    {
        "title": "Đăng ký tài khoản và xác thực email bằng OTP",
        "goal": "Tạo tài khoản học viên có email được xác thực và giảm đăng ký giả mạo.",
        "actors": "Khách chưa có tài khoản.",
        "entry": "/auth/register và /api/auth/register, /verify-registration-otp, /resend-registration-otp",
        "flow": [
            "Nhập tên, email, mật khẩu và xác nhận mật khẩu.",
            "Máy chủ chuẩn hóa email, kiểm tra trùng, kiểm tra độ mạnh mật khẩu rồi tạo/cập nhật user ở PENDING_VERIFICATION.",
            "Hệ thống sinh OTP 6 số, lưu bản băm, gửi email và trả thời điểm hết hạn/gửi lại.",
            "Nhập OTP đúng thì chuyển tài khoản sang ACTIVE; sai quá số lần quy định thì bị chặn.",
        ],
        "mechanism": [
            "Mật khẩu được băm bằng scrypt với salt ngẫu nhiên 16 byte; không lưu mật khẩu thô.",
            "OTP đăng ký hết hạn sau 10 phút, thời gian chờ gửi lại 60 giây và tối đa 5 lần nhập sai.",
            "Giới hạn yêu cầu dựa trên email, IP và dấu vân tay thiết bị; sự kiện bảo mật được lưu để audit.",
            "Đăng ký lại email chưa kích hoạt sẽ cập nhật hồ sơ pending thay vì tạo nhiều tài khoản.",
        ],
        "rules": [
            "Mật khẩu tối thiểu 8 ký tự, có chữ hoa, chữ thường, số và ký tự đặc biệt.",
            "Email ACTIVE không được đăng ký trùng.",
            "Không cấp phiên đăng nhập trước khi xác thực OTP.",
        ],
        "sources": ["app/api/auth/register/route.ts", "lib/registration-otp.ts", "lib/auth.ts", "prisma/schema.prisma"],
        "defense": "Điểm chính là tách hai bước tạo danh tính và kích hoạt danh tính. Dữ liệu tài khoản có thể được tạo ở trạng thái chờ, nhưng mọi quyền truy cập chỉ xuất hiện sau khi OTP hợp lệ.",
    },
    {
        "title": "Đăng nhập, phân quyền và thu hồi phiên",
        "goal": "Xác thực người dùng an toàn và áp dụng quyền STUDENT, TEACHER, ADMIN xuyên suốt giao diện/API.",
        "actors": "Tất cả người dùng đã kích hoạt.",
        "entry": "/auth/login, /api/auth/login, /api/auth/logout, /api/auth/me",
        "flow": [
            "Người dùng gửi email/mật khẩu; máy chủ kiểm tra rate limit, trạng thái tài khoản và mật khẩu.",
            "Nếu hợp lệ, máy chủ phát hành token HMAC chứa userId, role, authVersion và hạn dùng.",
            "Token được lưu trong cookie HttpOnly, SameSite=Lax, Secure ở production, thời hạn 7 ngày.",
            "Mỗi request đọc lại user từ PostgreSQL; nếu bị khóa, đổi role, accountStatus không ACTIVE hoặc authVersion lệch thì xóa cookie.",
        ],
        "mechanism": [
            "Chữ ký HMAC-SHA256 và timingSafeEqual chống sửa token/chênh lệch thời gian so sánh.",
            "requireUser/requireRole là cổng bảo vệ dùng lại ở page và API.",
            "Đổi mật khẩu, khóa tài khoản hoặc đổi role tăng authVersion và xóa Session cũ.",
            "Login rate limit: 30 lần/IP và 8 lần/email trong 15 phút.",
        ],
        "rules": [
            "Tài khoản chưa xác thực hoặc bị khóa không được đăng nhập.",
            "Admin không được tự khóa; không được khóa/hạ quyền admin hoạt động cuối cùng.",
            "Khi khóa/hạ quyền giáo viên, khóa các khóa ACTIVE do họ phụ trách.",
        ],
        "sources": ["lib/auth.ts", "lib/rate-limit.ts", "app/api/auth/login/route.ts", "app/api/admin/users/[userId]/route.ts"],
        "defense": "Token không được tin tuyệt đối sau khi phát hành. Hệ thống vẫn đối chiếu trạng thái user và authVersion ở database, nhờ đó quyền có thể bị thu hồi ngay mà không cần chờ token hết hạn.",
    },
    {
        "title": "Quên mật khẩu và hồ sơ cá nhân",
        "goal": "Cho phép khôi phục tài khoản, cập nhật thông tin và thay đổi mật khẩu có kiểm soát.",
        "actors": "Người dùng đã đăng ký.",
        "entry": "/auth/forgot-password, /profile, /api/profile, /api/profile/password",
        "flow": [
            "Quên mật khẩu: yêu cầu OTP qua email, nhập OTP và mật khẩu mới.",
            "Hồ sơ: xem/sửa tên, số điện thoại; email chỉ đọc ở form hiện tại.",
            "Đổi mật khẩu: xác minh mật khẩu hiện tại rồi ghi bản băm mới.",
            "Sau reset/đổi mật khẩu, vô hiệu hóa phiên cũ bằng authVersion và xóa session.",
        ],
        "mechanism": [
            "API yêu cầu OTP trả thông điệp trung tính để không làm lộ email có tồn tại.",
            "OTP reset lưu codeHash, expiresAt, attempts và consumedAt.",
            "Mật khẩu mới đi qua cùng bộ kiểm tra độ mạnh và scrypt như đăng ký.",
        ],
        "rules": [
            "Giới hạn gửi OTP theo email/IP và số lần nhập sai.",
            "OTP chỉ dùng một lần.",
            "Các thay đổi bảo mật quan trọng phải thu hồi quyền truy cập cũ.",
        ],
        "sources": ["app/api/auth/forgot-password/request-otp/route.ts", "app/api/auth/forgot-password/reset/route.ts", "app/api/profile/route.ts", "app/api/profile/password/route.ts"],
        "defense": "Luồng reset được thiết kế chống dò tài khoản và chống tái sử dụng OTP; thay mật khẩu đồng thời làm token cũ mất hiệu lực nhờ authVersion.",
    },
    {
        "title": "Thông báo và nhật ký email",
        "goal": "Đảm bảo sự kiện nghiệp vụ quan trọng đến đúng người và có dấu vết kiểm tra.",
        "actors": "Học viên, giảng viên, admin.",
        "entry": "Header/toast, /api/notifications, /api/teacher/revenue-notifications",
        "flow": [
            "Nghiệp vụ tạo Notification trong cùng transaction hoặc ngay sau khi thay đổi trạng thái.",
            "Header lấy thông báo chưa đọc, hiển thị toast và cho phép đánh dấu đã đọc.",
            "Khu vực doanh thu có feed và bộ đếm riêng cho giảng viên.",
            "Email OTP, duyệt hồ sơ, chứng nhận và thông báo hệ thống được ghi EmailLog SENT/FAILED.",
        ],
        "mechanism": [
            "Notification gắn userId, title, body, readAt và createdAt.",
            "Nodemailer/SMTP thực hiện gửi; lỗi email được ghi nhưng không rollback nghiệp vụ đã hoàn tất.",
            "Các luồng review/refund/report/withdrawal sinh thông báo theo trạng thái.",
        ],
        "rules": [
            "Người dùng chỉ đọc/đánh dấu thông báo của chính mình.",
            "Email thất bại cần lưu nội dung lỗi để admin quan sát.",
            "Thông báo tài chính được phân trang để không tải toàn bộ lịch sử.",
        ],
        "sources": ["app/api/notifications/route.ts", "lib/mailer.ts", "prisma/schema.prisma"],
        "defense": "Thông báo là kết quả quan sát được của workflow, còn EmailLog là bằng chứng vận hành. Việc tách hai lớp giúp hệ thống vẫn hoàn tất nghiệp vụ khi SMTP tạm lỗi.",
    },
    {
        "title": "Dashboard và khóa học của học viên",
        "goal": "Tổng hợp trạng thái học tập và đưa người học quay lại đúng điểm tiếp theo.",
        "actors": "Học viên; giáo viên/admin có thể học thử trong phạm vi cho phép.",
        "entry": "/student, /my-courses, /student/hoc-bai",
        "flow": [
            "Dashboard tổng hợp khóa đang học, khóa hoàn thành, điểm AI, bài test gần đây và hoạt động 7 ngày.",
            "Khóa học của tôi phân loại đang học/hoàn thành và hiển thị phần trăm tiến độ.",
            "Nút tiếp tục học dùng trạng thái gate để chọn lesson/test tiếp theo.",
            "Nếu enrollment ở REFUND_PENDING, giao diện và API đều chặn quyền học.",
        ],
        "mechanism": [
            "Enrollment xác nhận quyền truy cập; Feedback marker và LearningActivity cung cấp tiến độ/biểu đồ.",
            "getCourseLearningGateState hợp nhất lesson hoàn thành và test đã đạt.",
            "Dữ liệu gợi ý loại bỏ các khóa đã ghi danh.",
        ],
        "rules": [
            "Không chỉ dựa vào ẩn nút ở client; API học/test kiểm tra quyền lại.",
            "Tiến độ là số lesson hoàn thành trên tổng lesson.",
            "Hoàn thành khóa đòi hỏi toàn bộ module và test bắt buộc, không chỉ đạt 100% lesson.",
        ],
        "sources": ["app/student/page.tsx", "app/my-courses/page.tsx", "lib/course-learning-gates.ts", "lib/learning-progress.ts"],
        "defense": "Dashboard là lớp tổng hợp, còn quyết định học tiếp nằm ở gate phía server. Nhờ vậy người dùng tải lại trang hay gọi API trực tiếp vẫn không vượt được thứ tự học.",
    },
    {
        "title": "Học bài tuần tự, tiến độ đọc và chống tua video",
        "goal": "Xác minh người học thực sự đi qua nội dung và tuân thủ thứ tự module/test.",
        "actors": "Người đã ghi danh, chủ khóa, admin.",
        "entry": "/student/hoc-bai và /api/learning/lessons/[lessonId]/{start,heartbeat,complete}",
        "flow": [
            "Bắt đầu lesson: xác minh quyền, module đã mở rồi ghi mốc bắt đầu.",
            "Lesson văn bản yêu cầu tối thiểu 180 giây trước khi complete.",
            "Lesson video gửi heartbeat gồm vị trí/thời lượng; server cập nhật VideoWatchProgress.",
            "Complete video chỉ thành công khi heartbeat mới, không có seekViolation, đến gần cuối và thời gian thực xem tối thiểu 90%.",
            "Sau mỗi lesson/test đạt, hệ thống trả nextAction để chuyển đến nội dung tiếp theo.",
        ],
        "mechanism": [
            "Gate mở module đầu tiên; module sau chỉ mở khi module trước đủ lesson và test.",
            "Heartbeat cho phép vị trí tăng theo thời gian trôi qua; nhảy vượt ngưỡng hoặc đổi duration bất thường sẽ đánh dấu vi phạm.",
            "VideoWatchProgress lưu startedAt, lastHeartbeatAt, lastPositionSeconds, durationSeconds, seekViolation và completedAt.",
            "LearningActivity dùng sourceKey duy nhất để chống ghi trùng.",
        ],
        "rules": [
            "Heartbeat video phải mới trong 15 giây.",
            "Vị trí cuối phải cách cuối video không quá 2 giây.",
            "Giáo viên sở hữu và admin được preview nhưng học viên phải có enrollment ACTIVE.",
        ],
        "sources": ["lib/course-learning-gate-rules.ts", "app/api/learning/lessons/[lessonId]/heartbeat/route.ts", "app/api/learning/lessons/[lessonId]/complete/route.ts", "prisma/schema.prisma"],
        "defense": "Client chỉ báo vị trí video; server không tin tuyệt đối. Nó so vị trí với thời gian giữa hai heartbeat, lưu cờ vi phạm và kiểm tra nhiều điều kiện độc lập trước khi ghi hoàn thành.",
    },
    {
        "title": "Mua khóa học qua VNPay và chia doanh thu",
        "goal": "Ghi danh khóa trả phí an toàn, tránh xử lý trùng và phân bổ doanh thu minh bạch.",
        "actors": "Người mua, VNPay, giảng viên, admin.",
        "entry": "/api/courses/[id]/enroll, /api/payments/vnpay-return, /api/payments/vnpay-ipn",
        "flow": [
            "Khóa miễn phí hoặc khóa của chính giảng viên được enroll ngay.",
            "Khóa trả phí tạo Order + Payment PENDING và txnRef duy nhất, hạn thanh toán 15 phút.",
            "Server ký tham số VNPay; người dùng thanh toán trên cổng.",
            "Return/IPN xác minh chữ ký, mã giao dịch, trạng thái, số tiền và thời hạn.",
            "Thanh toán thành công tạo/upsert OrderItem và Enrollment.",
        ],
        "mechanism": [
            "Payment được claim PENDING → PROCESSING bằng updateMany để chặn hai callback xử lý cùng lúc.",
            "Các trạng thái cuối: PAID, FAILED, CANCELLED, EXPIRED; callback lặp trả ALREADY_PAID/ALREADY_FINAL.",
            "Khóa do TEACHER sở hữu chia 70% teacher, 30% admin; khóa khác 100% admin.",
            "Kiểm tra giá khóa hiện tại phải trùng số tiền Payment trước khi ghi nhận.",
        ],
        "rules": [
            "Không enroll trùng nhờ unique(userId, courseId) và upsert.",
            "Không tin dữ liệu redirect từ trình duyệt nếu chữ ký VNPay sai.",
            "Ví VND nội bộ đã ngừng; /api/wallet và /top-up trả HTTP 410.",
        ],
        "sources": ["app/api/courses/[id]/enroll/route.ts", "lib/course-payment.ts", "lib/vnpay.ts", "lib/revenue.ts"],
        "defense": "Tính đúng đắn tài chính dựa trên chữ ký, kiểm tra amount và xử lý idempotent. Dù VNPay gọi return và IPN gần như đồng thời, chỉ một tiến trình được quyền chuyển Payment khỏi PENDING.",
    },
    {
        "title": "Hoàn tiền khóa học",
        "goal": "Cho học viên yêu cầu hoàn tiền có điều kiện, đồng thời bảo vệ quyền học và doanh thu.",
        "actors": "Học viên/người mua và admin.",
        "entry": "/my-courses, /api/course-refunds, /api/admin/course-refunds",
        "flow": [
            "Người mua nhập lý do; server tìm enrollment và OrderItem gần nhất.",
            "Kiểm tra trong 7 ngày, tiến độ không quá 50%, khóa có phí và chưa từng có yêu cầu.",
            "Tạo CourseRefundRequest PENDING, chuyển Enrollment sang REFUND_PENDING và thông báo admin.",
            "Admin duyệt: xóa enrollment và dữ liệu học liên quan; từ chối: mở lại accessStatus ACTIVE.",
            "Tiền được hoàn ngoài hệ thống và kết quả gửi Notification.",
        ],
        "mechanism": [
            "Một OrderItem chỉ có tối đa một refundRequest.",
            "Tiến độ được tính từ marker PROGRESS so với tổng lesson.",
            "Refund được đưa vào bộ lọc doanh thu để không tính phần đã hoàn cho giảng viên.",
        ],
        "rules": [
            "Lý do 10–500 ký tự.",
            "Duyệt hoàn tiền xóa progress, attempts, AI assessments, activities, point transactions theo khóa và video progress.",
            "Refund hiện là EXTERNAL_ACCOUNT, chưa tự gọi API ngân hàng.",
        ],
        "sources": ["app/api/course-refunds/route.ts", "app/api/admin/course-refunds/[refundId]/route.ts", "lib/teacher-revenue.ts"],
        "defense": "Ngay khi gửi yêu cầu, quyền học bị khóa tạm thời để tránh vừa chờ hoàn vừa tiếp tục tiêu thụ nội dung. Khi admin quyết định, dữ liệu và quyền được xử lý nhất quán trong transaction.",
    },
    {
        "title": "Trung tâm bài test và điều kiện mở đề",
        "goal": "Cung cấp test cuối khóa, test theo module và bài luyện công khai theo một mô hình thống nhất.",
        "actors": "Học viên, giảng viên preview, admin.",
        "entry": "/student/tests, /student/tests/[testId], các API /student/tests/**",
        "flow": [
            "Danh sách hiển thị COURSE và PUBLIC_PRACTICE phù hợp.",
            "Khi mở test, server kiểm tra đề đủ 100 điểm, vai trò, enrollment và learning gate.",
            "Server tạo attemptToken ký HMAC với thời điểm bắt đầu/hết hạn.",
            "Client đếm ngược và tự nộp; server vẫn là nơi quyết định token còn hợp lệ.",
        ],
        "mechanism": [
            "TestKind: COURSE, PUBLIC_PRACTICE, TEACHER_ENTRANCE.",
            "AssessmentMode: STANDARD, WRITING, SPEAKING; QuestionType gồm trắc nghiệm, đúng/sai, điền, tự luận, nói.",
            "Module test mở sau khi lesson trong module hoàn tất; final test mở sau toàn bộ module/test trước.",
            "Lịch sử giữ mọi lượt làm; cấu hình hiện dùng số lượt không giới hạn nhưng vẫn có attemptNo.",
        ],
        "rules": [
            "Tổng score câu hỏi phải bằng đúng 100.",
            "Teacher entrance không được gọi từ luồng test học viên.",
            "Enrollment REFUND_PENDING bị chặn cả khi tải đề và khi nộp.",
        ],
        "sources": ["lib/student-test-data.ts", "lib/test-rules.ts", "lib/test-attempt-token.ts", "lib/course-learning-gates.ts"],
        "defense": "Điều kiện mở test được kiểm tra hai lần—lúc lấy đề và lúc nộp—để chống người dùng giữ trang cũ rồi vượt qua thay đổi quyền hoặc thứ tự học.",
    },
    {
        "title": "Chấm bài test, lưu kết quả và chứng nhận",
        "goal": "Chấm nhiều dạng câu hỏi, lưu bằng chứng lượt thi và hoàn tất khóa học đúng điều kiện.",
        "actors": "Người làm test; AI đối với câu tự luận/nói.",
        "entry": "/api/student/tests/[testId]/submit và trang result/history",
        "flow": [
            "Xác minh attemptToken và quyền làm bài.",
            "Câu khách quan so đáp án; câu essay/speaking gọi AI, quy đổi theo trọng số câu.",
            "Chuẩn hóa tổng điểm về thang 100 và so với passingScore.",
            "Lưu TestAttempt gồm answer snapshot, question results, điểm, trạng thái đạt và thời gian.",
            "Nếu đạt test COURSE cuối cùng, ghi hoàn thành và gửi email chứng nhận một lần.",
        ],
        "mechanism": [
            "Câu điền được NFKC, trim, gộp khoảng trắng và so không phân biệt hoa thường.",
            "AI lỗi/response không hợp lệ thì không lưu attempt, trả 502/503 để người dùng thử lại.",
            "Transaction Serializable gộp lưu attempt và trừ hạt đậu khi mua nhận xét chi tiết.",
            "CERT_SENT marker chống gửi email chứng nhận lặp.",
        ],
        "rules": [
            "Chế độ chỉ chấm điểm không trả feedback/mẫu chi tiết.",
            "Preview của chủ khóa không tạo TestAttempt thật.",
            "Course hoàn thành chỉ khi gateState.courseComplete, không chỉ vì một test đơn lẻ đạt.",
        ],
        "sources": ["app/api/student/tests/[testId]/submit/route.ts", "lib/test-ai-evaluation.ts", "lib/learning-progress.ts", "lib/mailer.ts"],
        "defense": "Mỗi TestAttempt là một snapshot có thể audit. Điểm AI được quy đổi theo trọng số từng câu, còn hoàn thành khóa dựa trên toàn bộ learning gate nên không thể đạt một test rồi bỏ qua module.",
    },
    {
        "title": "Lịch sử và phân tích kết quả học tập",
        "goal": "Giúp học viên xem lại bài làm và theo dõi xu hướng tiến bộ.",
        "actors": "Chủ kết quả; admin/giảng viên chỉ trong phạm vi được phép.",
        "entry": "/student/results, /student/results/[resultId], /student/tests/history",
        "flow": [
            "Hợp nhất TestAttempt, Writing AiAssessment và Speaking AiAssessment.",
            "Lọc theo loại, chuẩn hóa phần trăm và nhóm theo thời gian.",
            "Trang chi tiết hiển thị bài nộp, tiêu chí, lỗi, điểm mạnh/yếu, cải thiện và mẫu nếu đã mua.",
            "Lịch sử test lọc theo khóa/đề và cho phép làm lại.",
        ],
        "mechanism": [
            "student-results ánh xạ các nguồn dữ liệu khác nhau về một kiểu Result chung.",
            "Điểm thang 9/10/100 được chuẩn hóa để vẽ biểu đồ so sánh.",
            "Kiểm tra viewer trước khi trả chi tiết để chống IDOR.",
        ],
        "rules": [
            "Chỉ trả phần feedback mà người dùng đã mua/được phép xem.",
            "Mức IELTS/HSK/JLPT/TOPIK là tham chiếu nội bộ, không phải chứng chỉ chính thức.",
        ],
        "sources": ["lib/student-results.ts", "lib/student-test-attempt-result.ts", "app/student/results/ResultsClient.tsx"],
        "defense": "Giá trị của phần kết quả không chỉ là điểm cuối. Hệ thống chuẩn hóa nhiều thang điểm, giữ bài nộp và tách feedback chi tiết để người học có thể theo dõi tiến bộ qua thời gian.",
    },
    {
        "title": "Đánh giá khóa học",
        "goal": "Thu thập phản hồi đáng tin cậy từ người đã thực sự hoàn thành yêu cầu khóa.",
        "actors": "Học viên đã đạt test khóa học; người xem công khai.",
        "entry": "/api/courses/[id]/reviews và trang chi tiết khóa",
        "flow": [
            "Sau khi đạt khóa, người học mở form rating/comment.",
            "Server kiểm tra quyền review rồi upsert đánh giá của user cho course.",
            "Trang chi tiết tổng hợp điểm trung bình, số lượng và danh sách review.",
        ],
        "mechanism": [
            "Review hiện lưu trong Feedback dưới JSON có rating/comment.",
            "canReviewCourse kiểm tra TestAttempt đạt của test COURSE.",
            "Upsert logic cho phép chỉnh sửa thay vì tạo review trùng.",
        ],
        "rules": [
            "Rating 1–5; comment tối đa 1.000 ký tự.",
            "Người chưa đạt bài test không được đánh giá.",
            "Feedback đang overloaded với marker tiến độ; đây là điểm cần tách bảng khi mở rộng.",
        ],
        "sources": ["lib/course-reviews.ts", "app/api/courses/[id]/reviews/route.ts", "prisma/schema.prisma"],
        "defense": "Hệ thống ưu tiên review có xác thực trải nghiệm: chỉ người vượt qua test khóa mới được đánh giá, nhờ đó điểm sao phản ánh người học thật thay vì khách vãng lai.",
    },
    {
        "title": "Writing AI",
        "goal": "Tạo đề và chấm bài viết đa ngôn ngữ theo hai mức: điểm nhanh hoặc phản hồi chi tiết.",
        "actors": "Student, Teacher; Admin miễn phí.",
        "entry": "/student/writing-ai, /api/ai/writing-prompt, /api/ai/essay-evaluation",
        "flow": [
            "Chọn ngôn ngữ/task/chủ đề hoặc để AI sinh đề; có fallback khi tạo đề thất bại.",
            "Gửi bài viết; server validate độ dài/nội dung và quyền khóa học nếu có courseId.",
            "English dùng IELTS Writing; Trung/Nhật/Hàn dùng rubric chứng chỉ tương ứng.",
            "Kết quả được lưu AiAssessment và LearningActivity; feedback chi tiết mới trừ hạt đậu.",
        ],
        "mechanism": [
            "Ollama chat yêu cầu JSON, temperature 0, retry và timeout; parser/normalizer sửa cấu trúc phản hồi.",
            "Chấm IELTS theo tiêu chí Task Achievement/Response, Coherence, Lexical, Grammar; có kiểm tra task relevance/off-topic.",
            "Score-only loại bỏ nhận xét dài và model answer; detailed mode lưu feedback, mistakes, improvements, sampleAnswer.",
            "Lưu assessment và trừ điểm trong cùng transaction Serializable.",
        ],
        "rules": [
            "Writing feedback chi tiết hiện có giá 2 hạt đậu cho Student/Teacher; Admin không bị trừ.",
            "Nếu AI lỗi thì không trừ điểm và không lưu kết quả giả.",
            "Điểm English theo thang IELTS 9; ngôn ngữ khác theo thang 10.",
        ],
        "sources": ["app/api/ai/writing-prompt/route.ts", "app/api/ai/essay-evaluation/route.ts", "lib/ielts-grading.ts", "lib/ai-points.ts"],
        "defense": "Thiết kế tách 'chấm điểm' và 'phản hồi'. Người học vẫn nhận được điểm cơ bản miễn phí; hạt đậu chỉ được trừ khi yêu cầu nội dung AI chi tiết và việc trừ diễn ra cùng transaction với bản ghi kết quả.",
    },
    {
        "title": "Speaking AI",
        "goal": "Cho phép luyện nói, lưu bản ghi và chấm đa ngôn ngữ theo rubric.",
        "actors": "Student, Teacher; Admin miễn phí.",
        "entry": "/student/speaking-ai, Web Worker và /api/ai/speaking-evaluation/**",
        "flow": [
            "Chọn ngôn ngữ/task/topic hoặc sinh topic bằng AI.",
            "Trình duyệt xin microphone, ghi âm và dùng Web Worker để tạo transcript.",
            "Gửi transcript, audio, thời lượng và ngữ cảnh lên server.",
            "Server lưu audio, chấm transcript, lưu AiAssessment và trả điểm/feedback.",
        ],
        "mechanism": [
            "File audio được kiểm tra MIME + chữ ký nhị phân rồi lưu public/uploads/speaking.",
            "English dùng IELTS Speaking; Trung dùng HSK/HSKK; Nhật/Hàn dùng rubric tương ứng.",
            "Tiêu chí gồm fluency/coherence, vocabulary, grammar, pronunciation estimate và task relevance.",
            "Cấu hình loại bài/thời lượng đọc từ SystemSetting; duration gửi lên bị clamp theo cấu hình.",
        ],
        "rules": [
            "Speaking feedback chi tiết giá 7 hạt đậu cho Student/Teacher; Admin miễn phí.",
            "Cần transcript hợp lệ; AI/Ollama không khả dụng trả 503.",
            "Hiện evaluationBasis là TRANSCRIPT_ESTIMATE: audio được lưu để nghe lại nhưng chấm chủ yếu từ transcript, chưa phải phân tích âm vị trực tiếp.",
        ],
        "sources": ["app/student/speaking-ai/SpeakingAiClient.tsx", "public/workers/speaking-transcription.worker.mjs", "app/api/ai/speaking-evaluation/route.ts", "lib/speaking-ai-setting.ts"],
        "defense": "Điểm cần trình bày trung thực: hệ thống có pipeline ghi âm và lưu audio, nhưng điểm hiện tại chủ yếu suy ra từ transcript. Đây là phiên bản khả dụng và có đường nâng cấp rõ ràng sang ASR/acoustic scoring chuyên dụng.",
    },
    {
        "title": "Hạt đậu AI và mua điểm qua VNPay",
        "goal": "Tạo cơ chế trả phí theo lượt cho phản hồi AI mà không dùng ví tiền VND nội bộ.",
        "actors": "Student, Teacher, Admin.",
        "entry": "/student/wallet, /student/rewards, /api/ai/points/**",
        "flow": [
            "Người dùng chọn số hạt, tạo Payment PENDING và thanh toán VNPay.",
            "Callback thành công tạo PointTransaction AI_POINTS_PURCHASE.",
            "Khi dùng feedback, hệ thống kiểm tra số dư và ghi transaction âm.",
            "Trang phần thưởng hiển thị earned, spent, available, streak và lịch sử.",
        ],
        "mechanism": [
            "Số dư được tính từ ledger PointTransaction thay vì một biến dễ lệch.",
            "sourceKey unique bảo đảm cộng/trừ idempotent.",
            "Mua điểm kiểm tra expectedAmount = points × AI_POINT_PRICE_VND; mặc định 1.000 VND/hạt.",
            "Serializable transaction ngăn hai request đồng thời tiêu vượt số dư.",
        ],
        "rules": [
            "Chỉ AI_POINTS_PURCHASE và AI_POINTS_ADMIN_GRANT được tính là nguồn cộng hợp lệ.",
            "Student/Teacher bị tính phí; Admin miễn phí.",
            "Cơ chế thưởng streak/hoàn thành hiện chưa cấp điểm đáng kể.",
        ],
        "sources": ["lib/ai-points.ts", "app/api/ai/points/buy/route.ts", "app/api/ai/points/vnpay-ipn/route.ts", "app/student/rewards/page.tsx"],
        "defense": "Hạt đậu là sổ cái append-only. Mỗi giao dịch có sourceKey duy nhất và balanceAfter để audit, nên callback lặp hoặc request lặp không thể cộng/trừ hai lần.",
    },
    {
        "title": "Đăng ký giảng viên và hồ sơ chứng chỉ",
        "goal": "Xây dựng workflow nâng quyền có xét duyệt dựa trên ngôn ngữ, chứng chỉ và bài thi đầu vào.",
        "actors": "Người nộp hồ sơ, admin.",
        "entry": "/teacher-registration và /api/teacher-applications/**",
        "flow": [
            "Admin bật/tắt nhận hồ sơ; người dùng chọn ngôn ngữ và tải 1–3 chứng chỉ.",
            "Nếu ngôn ngữ không có entrance test, hồ sơ chuyển chờ review.",
            "Nếu có test, hồ sơ ở DRAFT, trải qua phiên thi tuần tự rồi SUBMITTED/UNDER_REVIEW.",
            "Admin duyệt chuyển user thành TEACHER; từ chối lưu lý do và gửi notification/email.",
        ],
        "mechanism": [
            "TeacherApplication giữ attemptNo, trạng thái, entrance test/attempt, reviewer và log.",
            "TeacherCertificate lưu đường dẫn, tên file, loại và ngày hết hạn.",
            "TeacherApplicationLog tạo lịch sử trạng thái cho từng lần nộp.",
        ],
        "rules": [
            "Chứng chỉ JPG/PNG/PDF, tối đa 10 MB/file, bắt buộc ngày hết hạn.",
            "Giáo viên chỉ được tạo khóa theo ngôn ngữ đã được phê duyệt.",
            "Các trạng thái gồm DRAFT, SUBMITTED, UNDER_REVIEW, APPROVED, REJECTED, EXPIRED, FAILED_CHEATING.",
        ],
        "sources": ["app/api/teacher-applications/route.ts", "lib/teacher-registration-data.ts", "lib/teacher-onboarding.ts", "prisma/schema.prisma"],
        "defense": "Nâng role không diễn ra trực tiếp từ form. Hồ sơ, chứng chỉ, bài thi và quyết định admin được liên kết thành một workflow có log, nhờ đó có thể giải thích vì sao một tài khoản trở thành giảng viên.",
    },
    {
        "title": "Bài thi đầu vào tuần tự và anti-cheat",
        "goal": "Giám sát phiên thi giảng viên và hạn chế rời bài, nhiều phiên, tua thời gian hoặc lộ câu hỏi.",
        "actors": "Ứng viên giảng viên, admin review.",
        "entry": "TeacherSequentialExam, TeacherAntiCheatOverlays và các API heartbeat/session/question",
        "flow": [
            "Ứng viên đồng ý quy định, bật camera, vào fullscreen rồi server mới ghi nhận bắt đầu.",
            "Câu hỏi được snapshot và khóa; người thi chủ động reveal từng câu, nhận transition token và đồng hồ riêng.",
            "Client gửi heartbeat mỗi 3 giây với sessionId, sequence, camera, fullscreen, focus, visibility và extended display.",
            "Rời fullscreen/tab/focus tạo incident; incidentId khử trùng nhiều tín hiệu của cùng một hành động.",
            "Đủ 3 violation đếm được thì hồ sơ FAILED_CHEATING và tạo TestAttempt điểm 0.",
        ],
        "mechanism": [
            "TeacherEntranceQuestionInstance giữ snapshot nội dung/đáp án, thứ tự, hash, trạng thái, thời điểm reveal/deadline và token băm.",
            "Speaking mặc định 60 giây chuẩn bị + 120 giây trả lời; Writing mặc định 60 phút; deadline được kiểm tra server-side.",
            "Heartbeat gap, camera off, nhiều màn hình/phiên là warning; TAB_HIDDEN, WINDOW_BLUR, FULLSCREEN_EXIT là violation đếm.",
            "SuspiciousEvent tổng hợp theo eventType; AntiCheatLog giữ chi tiết và metadata.",
        ],
        "rules": [
            "Website không thể khóa Alt+Tab/Esc tuyệt đối; hệ thống phát hiện và ghi nhận thay vì tuyên bố ngăn hoàn toàn.",
            "Heartbeat quá hạn có thể chặn nộp; tải lại không đặt lại đồng hồ câu hỏi.",
            "Sau khi nộp thành công, client thoát fullscreen, dừng camera/heartbeat và ngừng ghi sự kiện.",
        ],
        "sources": ["lib/teacher-anti-cheat.ts", "lib/teacher-sequential-exam.ts", "app/api/teacher-applications/[applicationId]/heartbeat/route.ts", "app/teacher-registration/components/TeacherAntiCheatOverlays.tsx"],
        "defense": "Anti-cheat không dựa vào một tín hiệu duy nhất. Nó kết hợp fullscreen, visibility, focus, camera, phiên heartbeat và snapshot thời gian phía server; đồng thời dùng incidentId để không phạt ba lần cho cùng một lần Alt+Tab.",
    },
    {
        "title": "Dashboard giảng viên",
        "goal": "Cho giảng viên theo dõi khóa học, học viên, doanh thu, test và báo cáo cần xử lý.",
        "actors": "Giảng viên.",
        "entry": "/teacher",
        "flow": [
            "Tải số khóa, học viên, doanh số, doanh thu khả dụng/đang giữ và trạng thái báo cáo.",
            "Hiển thị khóa gần đây, biểu đồ báo cáo và danh sách vấn đề mới.",
            "Điều hướng nhanh đến course, student, test, report và revenue.",
        ],
        "mechanism": [
            "Dữ liệu luôn lọc theo instructorId của user hiện tại.",
            "Doanh thu dựa trên OrderItem hợp lệ; báo cáo dựa trên CourseReport của khóa sở hữu.",
            "Khu dashboard chỉ tổng hợp, thao tác ghi được chuyển sang API chuyên trách.",
        ],
        "rules": [
            "Không cho giáo viên xem dữ liệu khóa của người khác.",
            "Số doanh thu phải loại refund APPROVED và trừ withdrawal đang giữ.",
        ],
        "sources": ["app/teacher/page.tsx", "lib/teacher-revenue.ts", "app/api/course-reports/route.ts"],
        "defense": "Dashboard giảng viên là bảng điều khiển theo ownership: mọi thống kê đều có điều kiện instructorId, tránh biến dashboard thành điểm rò rỉ dữ liệu giữa các giảng viên.",
    },
    {
        "title": "Quản lý khóa học và quy trình duyệt",
        "goal": "Cho giảng viên xây khóa, còn admin kiểm soát chất lượng trước khi công khai/xóa.",
        "actors": "Teacher, Admin.",
        "entry": "/teacher/courses và /api/teacher/courses/**, /api/admin/course-approval",
        "flow": [
            "Tạo khóa ở PENDING_APPROVAL; cấu hình tên, mô tả, giá, ngôn ngữ, category, level, duration, thumbnail.",
            "Thêm module, lesson và test; khi sẵn sàng thì submitForApproval.",
            "Admin duyệt → ACTIVE hoặc từ chối → REJECTED; có thể LOCK/UNLOCK.",
            "Giáo viên yêu cầu xóa → PENDING_DELETE; admin duyệt hoặc phục hồi trạng thái trước đó.",
        ],
        "mechanism": [
            "getCourseReadiness yêu cầu ngôn ngữ, ít nhất một module, một lesson, test cuối khóa và mỗi test đủ 100 điểm.",
            "Auto-approval chỉ áp dụng khi khóa đã sẵn sàng; khóa mới luôn pending vì chưa có cấu trúc.",
            "Khóa có enrollment/order/payment không bị xóa vật lý mà chuyển LOCKED để giữ lịch sử tài chính.",
            "Khóa chưa phát sinh quan hệ có thể xóa cùng các dữ liệu con trong transaction.",
        ],
        "rules": [
            "Teacher chỉ dùng ngôn ngữ trong hồ sơ APPROVED; Admin dùng ngôn ngữ hệ thống active.",
            "Giá không âm; status do Admin kiểm soát.",
            "Chỉ Admin được khóa/mở khóa và quyết định xóa cuối cùng.",
        ],
        "sources": ["app/api/teacher/courses/route.ts", "app/api/teacher/courses/[courseId]/route.ts", "lib/course-readiness.ts", "lib/course-approval.ts"],
        "defense": "Quy trình duyệt dựa trên readiness thay vì chỉ nút bấm. Admin không thể mở ACTIVE nếu cấu trúc khóa thiếu module, lesson hoặc test đủ 100 điểm.",
    },
    {
        "title": "Quản lý module và lesson",
        "goal": "Tổ chức nội dung khóa theo chương/bài và hỗ trợ video.",
        "actors": "Chủ khóa và Admin.",
        "entry": "/teacher/courses/[courseId]/modules/[moduleId] và API modules/lessons",
        "flow": [
            "Tạo/sửa/xóa module và đặt thứ tự.",
            "Tạo/sửa/xóa lesson, nội dung và video; xem preview trước khi lưu.",
            "Danh sách học tập sắp xếp theo order và được gate sử dụng.",
        ],
        "mechanism": [
            "Mọi API tải course → kiểm tra instructorId hoặc ADMIN trước khi ghi.",
            "Video upload được kiểm tra loại/chữ ký file và lưu public local.",
            "Lesson thuộc Module, Module thuộc Course; xóa được kiểm soát theo quan hệ.",
        ],
        "rules": [
            "Không được chỉnh module/lesson của khóa người khác.",
            "Video hỗ trợ MP4/WebM/MOV theo giao diện, giới hạn 500 MB ở route upload.",
            "Thay đổi nội dung có thể đưa khóa về trạng thái cần duyệt tùy cấu hình.",
        ],
        "sources": ["app/api/teacher/courses/[courseId]/modules/route.ts", "app/api/teacher/courses/[courseId]/modules/[moduleId]/lessons/route.ts", "app/api/teacher/upload/route.ts"],
        "defense": "Module/lesson là cây dữ liệu có ownership rõ ràng. Mỗi thao tác ghi đều đi từ course để xác minh chủ sở hữu, thay vì tin moduleId/lessonId do client gửi.",
    },
    {
        "title": "Quản lý đề test và câu hỏi",
        "goal": "Cho giáo viên/admin tạo đề đủ điều kiện chấm và hỗ trợ tài liệu/audio.",
        "actors": "Teacher với khóa sở hữu; Admin với public/entrance test.",
        "entry": "/teacher/tests, /teacher/tests/[testId]/questions và API /teacher/tests/**",
        "flow": [
            "Tạo test với kind/mode, tên, mô tả, passingScore, timeLimit, shuffle.",
            "Thêm câu hỏi, đáp án, score, giải thích, hint, audio và thời gian riêng entrance.",
            "Theo dõi tổng điểm còn thiếu; chỉ sẵn sàng khi bằng 100.",
            "Upload tài liệu đề hoặc audio câu hỏi.",
        ],
        "mechanism": [
            "Teacher chỉ tạo COURSE test cho khóa của mình; Admin tạo PUBLIC_PRACTICE/TEACHER_ENTRANCE.",
            "Teacher entrance bắt buộc assessment mode Writing hoặc Speaking.",
            "Thời gian Speaking cho entrance giới hạn 30–300 giây trả lời, preparation 0–300; Writing 60–10.800 giây.",
            "Upload validation kiểm tra MIME và magic bytes, không chỉ phần mở rộng.",
        ],
        "rules": [
            "Mỗi khóa có cấu trúc test được ràng buộc theo nghiệp vụ; course phải có module trước khi tạo.",
            "Không cho tổng điểm vượt 100; publish/readiness yêu cầu đúng 100.",
            "Audio tối đa 100 MB, tài liệu test tối đa 50 MB.",
        ],
        "sources": ["app/api/teacher/tests/route.ts", "app/api/teacher/tests/[testId]/questions/route.ts", "lib/test-rules.ts", "lib/upload-validation.ts"],
        "defense": "Tính hợp lệ của đề được kiểm tra ở cả lúc thêm câu và lúc người học mở/nộp. Điều này ngăn một đề chưa đủ 100 điểm bị chấm sai chỉ vì giao diện đã cho lưu.",
    },
    {
        "title": "Quản lý học viên và tài khoản",
        "goal": "Cho giáo viên xem học viên của mình và admin quản trị tài khoản toàn hệ thống.",
        "actors": "Teacher, Admin.",
        "entry": "/teacher/students, /api/teacher/students, /api/admin/users/[userId]",
        "flow": [
            "Teacher tìm theo tên/email/khóa và xem enrollment của khóa mình.",
            "Admin xem toàn bộ user, đổi email/role và khóa/mở tài khoản.",
            "Thay đổi bảo mật làm tăng authVersion và thu hồi session.",
        ],
        "mechanism": [
            "Teacher query qua enrollment → course.instructorId.",
            "Admin update trong transaction; nếu giáo viên bị khóa/hạ role thì khóa các khóa ACTIVE.",
            "Bảo vệ admin cuối cùng bằng count other active admins.",
        ],
        "rules": [
            "Admin không thể tự khóa chính mình.",
            "Không được khóa/hạ quyền admin hoạt động cuối cùng.",
            "Không trả dữ liệu học viên nằm ngoài khóa giáo viên sở hữu.",
        ],
        "sources": ["lib/teacher-students.ts", "app/api/teacher/students/route.ts", "app/api/admin/users/[userId]/route.ts"],
        "defense": "Quản lý user gắn với tác động dây chuyền: khóa giáo viên không chỉ đổi một cờ, mà còn thu hồi phiên và khóa khóa học đang công khai để tránh nội dung không còn người chịu trách nhiệm.",
    },
    {
        "title": "Báo cáo khóa học",
        "goal": "Cho học viên phản ánh vấn đề và phân luồng xử lý giữa giáo viên/admin.",
        "actors": "Học viên đã ghi danh, Teacher sở hữu, Admin.",
        "entry": "CourseReportButton, /teacher/reports, /api/course-reports",
        "flow": [
            "Học viên chọn toàn khóa hoặc lesson, category, tiêu đề/nội dung và xác nhận trung thực.",
            "Server kiểm tra enrollment/lesson, chống spam rồi tạo report và thông báo admin.",
            "Giáo viên có thể claim/chuyển IN_REVIEW; Admin quyết định RESOLVED/REJECTED.",
            "Phản hồi/trạng thái mới tạo notification cho người báo cáo.",
        ],
        "mechanism": [
            "Phân quyền GET: Admin xem tất cả, Teacher xem khóa sở hữu, Student xem report của mình.",
            "Claim dùng updateMany có điều kiện respondedById để chống hai người cùng tiếp nhận.",
            "Search bao phủ title, description, response, course, lesson, reporter; có phân trang.",
        ],
        "rules": [
            "Tiêu đề 5–160, nội dung 20–5.000; tối đa 5 report/giờ.",
            "Teacher chỉ đưa về IN_REVIEW; trạng thái cuối do Admin.",
            "Report đã kết thúc không được nhận xử lý lại.",
        ],
        "sources": ["app/api/course-reports/route.ts", "app/api/course-reports/[reportId]/route.ts", "app/teacher/reports/page.tsx"],
        "defense": "Cơ chế claim ngăn hai người phản hồi chồng chéo. Giáo viên được tham gia xử lý nhưng quyền đóng/từ chối cuối cùng vẫn thuộc admin để giữ tính độc lập.",
    },
    {
        "title": "Doanh thu giảng viên và tài khoản ngân hàng",
        "goal": "Tính doanh thu khả dụng và xác minh tài khoản nhận tiền trước khi rút.",
        "actors": "Teacher, Admin.",
        "entry": "/teacher/revenue, /api/teacher/bank-account/**, /api/teacher/revenue-withdrawals",
        "flow": [
            "Teacher nhập ngân hàng, chi nhánh, số tài khoản, tên chủ tài khoản.",
            "Hệ thống gửi OTP email; xác minh thành công mới đặt bank account VERIFIED.",
            "Dashboard tính doanh thu teacher từ OrderItem, bỏ refund approved và trừ khoản đang giữ.",
            "Teacher tạo withdrawal với snapshot tài khoản đã xác minh.",
        ],
        "mechanism": [
            "Account name chuẩn hóa uppercase; account number 6–30 số; hiển thị dạng che.",
            "OTP tài khoản ngân hàng hết hạn 10 phút, tối đa 5 lần thử, có cooldown gửi lại.",
            "Lưu change log với giá trị cũ/mới đã che, IP và device fingerprint.",
            "Tạo yêu cầu rút trong transaction Serializable và tính lại available ngay trong transaction.",
        ],
        "rules": [
            "Số tiền rút là số nguyên dương và không vượt available.",
            "Các trạng thái PENDING, APPROVED, PAID, COMPLETED giữ tiền; REJECTED giải phóng.",
            "Yêu cầu lưu snapshot ngân hàng để thay đổi tài khoản sau này không làm sai lịch sử.",
        ],
        "sources": ["lib/teacher-bank-account.ts", "app/api/teacher/bank-account/request-otp/route.ts", "app/api/teacher/revenue-withdrawals/route.ts", "lib/teacher-revenue.ts"],
        "defense": "Điểm mạnh tài chính là tính lại số dư trong Serializable transaction và snapshot tài khoản nhận tiền. Hai request rút đồng thời không thể cùng nhìn thấy một số dư cũ rồi rút vượt.",
    },
    {
        "title": "Khiếu nại rút doanh thu",
        "goal": "Cho giảng viên khiếu nại sau khi hệ thống ghi nhận đã chuyển tiền.",
        "actors": "Teacher, Admin.",
        "entry": "/api/teacher/revenue-withdrawals/[withdrawalId]/complaints và API admin tương ứng",
        "flow": [
            "Chỉ withdrawal PAID/COMPLETED mới mở form khiếu nại.",
            "Teacher chọn chưa nhận, sai số tiền hoặc khác; có thể nhập số thực nhận và tải ảnh.",
            "Admin xem bằng chứng, ghi chú rồi RESOLVED hoặc REJECTED.",
            "Kết quả được thông báo lại cho teacher.",
        ],
        "mechanism": [
            "Quan hệ một-một giữa withdrawal và complaint chống nộp trùng.",
            "Ảnh evidence được kiểm tra và lưu local; metadata tên/URL được lưu cùng complaint.",
            "Admin route kiểm tra withdrawalId/complaintId khớp trước khi cập nhật.",
        ],
        "rules": [
            "Mỗi withdrawal tối đa một complaint.",
            "reportedAmount bắt buộc/hợp lệ khi lý do là sai số tiền.",
            "Chuyển khoản thực tế vẫn ngoài hệ thống.",
        ],
        "sources": ["app/api/teacher/revenue-withdrawals/[withdrawalId]/complaints/route.ts", "app/api/admin/revenue-withdrawals/[withdrawalId]/complaints/[complaintId]/route.ts", "lib/withdrawal-complaint-evidence.ts"],
        "defense": "Khiếu nại được gắn trực tiếp với withdrawal và snapshot ngân hàng, nên admin có đủ ngữ cảnh để đối soát thay vì xử lý một ticket rời rạc.",
    },
    {
        "title": "Quản trị hồ sơ, khóa học, hoàn tiền và rút tiền",
        "goal": "Cung cấp một trung tâm kiểm soát các workflow cần quyết định của con người.",
        "actors": "Admin.",
        "entry": "/admin và các API /api/admin/**",
        "flow": [
            "Duyệt/từ chối hồ sơ giảng viên và xem anti-cheat.",
            "Duyệt/khóa/mở/xóa khóa học theo readiness.",
            "Duyệt/từ chối refund và withdrawal; xác nhận PAID/COMPLETED.",
            "Xử lý complaint và course report.",
        ],
        "mechanism": [
            "Mỗi action kiểm tra role ADMIN ở server.",
            "Các cập nhật quan trọng dùng transaction và kiểm tra trạng thái hiện tại để chống xử lý hai lần.",
            "Notification/email được tạo sau quyết định; dashboard hiển thị bộ đếm pending.",
        ],
        "rules": [
            "Quyết định từ chối thường yêu cầu lý do/ghi chú.",
            "Không cho chuyển trạng thái từ trạng thái cuối không hợp lệ.",
            "Admin vẫn phải tuân thủ readiness và kiểm tra doanh thu, không có quyền bỏ qua toàn bộ quy tắc.",
        ],
        "sources": ["app/admin/AdminDashboard.tsx", "app/api/admin/teacher-applications/[applicationId]/review/route.ts", "app/api/admin/course-refunds/[refundId]/route.ts", "app/api/admin/revenue-withdrawals/[withdrawalId]/route.ts"],
        "defense": "Admin không phải là một nút 'sửa mọi thứ'. Mỗi tab là một state machine riêng, kiểm tra trạng thái nguồn và quy tắc trước khi cho chuyển trạng thái.",
    },
    {
        "title": "Cấu hình nghiệp vụ hệ thống",
        "goal": "Cho phép thay đổi một số hành vi vận hành mà không sửa mã nguồn.",
        "actors": "Admin.",
        "entry": "/api/admin/course-approval, /teacher-entrance, /speaking-config, /api/languages",
        "flow": [
            "Bật/tắt nhận hồ sơ giảng viên.",
            "Bật/tắt auto approval khóa khi khóa đã đủ readiness.",
            "Cấu hình loại bài/thời lượng Speaking AI.",
            "Thêm/kích hoạt ngôn ngữ học.",
        ],
        "mechanism": [
            "SystemSetting lưu key-value JSON; mỗi thư viện có key và kiểu đọc riêng.",
            "LearningLanguage là bảng danh mục dùng chung cho user, course, test, application.",
            "Khi bật lại đăng ký có thể gửi notification/email cho học viên.",
        ],
        "rules": [
            "Chỉ Admin ghi cấu hình.",
            "Speaking duration bị giới hạn 30–900 giây.",
            "Một số API cấu hình hiện chưa có UI admin hoàn chỉnh.",
        ],
        "sources": ["lib/course-approval.ts", "lib/teacher-onboarding.ts", "lib/speaking-ai-setting.ts", "app/api/languages/route.ts"],
        "defense": "Tách SystemSetting giúp thay đổi chính sách vận hành mà không deploy lại. Tuy nhiên giao diện quản trị còn chưa bao phủ hết API—đây là hạn chế được xác định rõ.",
    },
    {
        "title": "Analytics quản trị và xuất báo cáo",
        "goal": "Đo sức khỏe sản phẩm, tài chính, học tập, AI, bảo mật và hiệu quả giảng viên.",
        "actors": "Admin.",
        "entry": "/admin tab Analytics, /api/admin/analytics",
        "flow": [
            "Chọn preset hôm nay/hôm qua/7 ngày/30 ngày/tuần/tháng/quý/năm hoặc khoảng tùy chỉnh.",
            "Server truy vấn dữ liệu trong khoảng rồi tổng hợp series, KPI và ranking.",
            "UI hiển thị user growth, course/enrollment, revenue, test/AI, refund/withdrawal, report/security/email.",
            "Cho phép xuất CSV, XLSX và PDF.",
        ],
        "mechanism": [
            "admin-analytics chuẩn hóa timezone/range và bucket theo ngày/tuần/tháng/năm.",
            "Doanh thu tách course revenue, admin/teacher share và AI point revenue.",
            "Ranking gồm học viên, giảng viên, khóa theo nhiều tiêu chí.",
            "jsPDF, autotable và xlsx tạo file xuất ở client.",
        ],
        "rules": [
            "Chỉ Admin gọi API.",
            "Refund approved phải được phản ánh trong chỉ số hoàn tiền/doanh thu.",
            "Email SENT/FAILED và anti-cheat log được đưa vào quan sát vận hành.",
        ],
        "sources": ["lib/admin-analytics.ts", "app/admin/AnalyticsDashboard.tsx", "app/admin/AnalyticsDashboardSections.tsx"],
        "defense": "Analytics không chỉ đếm user. Nó nối vòng đời từ marketing → enrollment → học/test/AI → doanh thu/refund/withdrawal → report/security, giúp admin thấy cả tăng trưởng và rủi ro.",
    },
    {
        "title": "Upload file và kiểm tra nội dung",
        "goal": "Cho phép tải thumbnail, video, audio, tài liệu, chứng chỉ và bằng chứng với kiểm soát cơ bản.",
        "actors": "Teacher, ứng viên, học viên trong Speaking, Admin.",
        "entry": "Các API upload chuyên biệt.",
        "flow": [
            "Client gửi multipart/form-data.",
            "Server kiểm tra role/ownership, dung lượng, MIME, extension và magic bytes.",
            "Tạo tên file an toàn rồi ghi vào public/uploads theo từng loại.",
            "Lưu URL tương đối trong database.",
        ],
        "mechanism": [
            "upload-validation hỗ trợ image, PDF, audio, video và nhận diện chữ ký JPEG/PNG/PDF/MP3/WAV/Ogg/WebM/MP4…",
            "Giới hạn thay đổi theo ngữ cảnh: certificate 10 MB, audio 100 MB, test material 50 MB, video giao diện 500 MB.",
            "Tên file do server sinh, tránh dùng trực tiếp tên người dùng làm đường dẫn.",
        ],
        "rules": [
            "Không tin MIME do trình duyệt gửi; phải khớp magic bytes.",
            "Chỉ chủ tài nguyên/Admin được upload/chỉnh sửa liên quan.",
            "Lưu local phù hợp demo một máy nhưng cần object storage khi scale/serverless.",
        ],
        "sources": ["lib/upload-validation.ts", "app/api/teacher/upload/route.ts", "app/api/teacher/test-material-upload/route.ts", "app/api/teacher-applications/route.ts"],
        "defense": "Kiểm tra file có hai lớp: metadata và chữ ký nhị phân. Điểm cần nâng cấp là nơi lưu—public local không bền khi chạy nhiều instance, vì vậy roadmap là S3/R2 kèm signed URL và quét malware.",
    },
    {
        "title": "Health check, logging và khả năng phục hồi",
        "goal": "Phát hiện nhanh trạng thái ứng dụng/database và giữ dấu vết lỗi vận hành.",
        "actors": "DevOps/Admin/hệ thống giám sát.",
        "entry": "/api/health, console logs, EmailLog và các bảng audit.",
        "flow": [
            "Health endpoint chạy SELECT 1 và trả latency.",
            "Database lỗi trả HTTP 503 thay vì báo healthy giả.",
            "AI có health check riêng qua Ollama /api/tags; lỗi AI trả 503/502 theo loại.",
            "Email, anti-cheat, application, bank change và payment raw response tạo audit trail.",
        ],
        "mechanism": [
            "PrismaPg adapter kết nối PostgreSQL; Prisma client được cache trong development.",
            "Các callback tài chính lưu responseCode, transactionStatus và rawResponse.",
            "Ollama có timeout, retry tăng dần và kiểm tra response bị cắt.",
        ],
        "rules": [
            "Không log secret/mật khẩu/OTP thô.",
            "Health database tách khỏi health AI; một dịch vụ phụ lỗi không nên báo sai trạng thái toàn bộ app.",
            "Production bắt buộc AUTH_SECRET mạnh và cấu hình SMTP/VNPay/Ollama đầy đủ.",
        ],
        "sources": ["app/api/health/route.ts", "lib/prisma.ts", "lib/ai/ollama-service.ts", "prisma/schema.prisma"],
        "defense": "Khả năng quan sát được xây từ endpoint health và các bảng audit theo nghiệp vụ. Hệ thống phân biệt lỗi AI với lỗi database để phản hồi đúng và tránh ghi dữ liệu nửa chừng.",
    },
    {
        "title": "Đa ngôn ngữ và định tuyến tương thích",
        "goal": "Hỗ trợ nội dung/giao diện theo ngôn ngữ và giữ các URL cũ không làm gãy trải nghiệm.",
        "actors": "Người học tiếng Anh, Trung, Nhật, Hàn và người dùng cũ.",
        "entry": "language-display, test-language-labels, speaking/writing-languages và các route redirect.",
        "flow": [
            "Course/Test/Application gắn LearningLanguage.",
            "UI lấy label theo mã ngôn ngữ; Speaking/Writing ánh xạ sang locale/rubric.",
            "Các route cũ /wallet, /my-learning, /student/lam-bai, /forgot-password redirect sang route mới.",
        ],
        "mechanism": [
            "Các thư viện label tập trung giảm hard-code trong component.",
            "Speech locale ánh xạ zh-CN, ja-JP, ko-KR, vi-VN, en-US.",
            "Redirect server-side giữ bookmark/link cũ hoạt động.",
        ],
        "rules": [
            "Không dùng dữ liệu ngôn ngữ không active cho nội dung mới.",
            "Teacher bị cố định ngôn ngữ theo application được duyệt.",
            "Một số chuỗi nguồn còn mojibake và cần chuẩn hóa UTF-8.",
        ],
        "sources": ["lib/language-display.ts", "lib/test-language-labels.ts", "lib/speaking-languages.ts", "app/wallet/page.tsx"],
        "defense": "Đa ngôn ngữ không chỉ là dịch nhãn; nó còn quyết định speech locale, rubric chấm và ngôn ngữ khóa mà giáo viên được phép dạy.",
    },
]


QA_ITEMS = [
    (
        "Vì sao dùng cookie tự ký thay vì thư viện auth?",
        "Mục tiêu của đồ án là minh họa cơ chế rõ ràng: HMAC-SHA256, HttpOnly, hạn 7 ngày và authVersion. Điểm hạn chế là tự vận hành nhiều trách nhiệm; production có thể chuyển sang Auth.js/OIDC nhưng vẫn giữ nguyên RBAC và kiểm tra trạng thái database.",
    ),
    (
        "Token bị đánh cắp thì thu hồi thế nào?",
        "Tăng authVersion khi đổi mật khẩu/role/khóa tài khoản. authenticate đọc user mỗi request và từ chối token có ver cũ, đồng thời xóa cookie.",
    ),
    (
        "Tại sao rate limit hiện chưa đủ khi scale?",
        "Bucket đang ở Map trong bộ nhớ process. Nhiều instance sẽ có bộ đếm riêng; production nên chuyển sang Redis/Upstash và cấu hình trusted proxy cho IP.",
    ),
    (
        "Làm sao chống callback VNPay chạy hai lần?",
        "Payment có txnRef unique và trạng thái. Callback đầu claim PENDING → PROCESSING; callback sau thấy đã final/đã paid. OrderItem và Enrollment còn dùng upsert/unique.",
    ),
    (
        "Vì sao dùng transaction Serializable khi trừ hạt đậu/rút tiền?",
        "Để hai request đồng thời không cùng đọc một số dư cũ rồi cùng chi tiêu. Transaction tính lại số dư và ghi ledger/yêu cầu trong một đơn vị nguyên tử.",
    ),
    (
        "Người học có thể gọi thẳng API complete lesson không?",
        "API complete tự kiểm tra enrollment, accessStatus, module gate và bằng chứng thời gian/video. Việc ẩn nút client không phải lớp bảo vệ chính.",
    ),
    (
        "Chống tua video có tuyệt đối không?",
        "Không tuyệt đối, nhưng có kiểm chứng server: heartbeat, vị trí cho phép theo thời gian, cờ seekViolation, vị trí cuối, 90% thời gian và heartbeat mới. Production có thể bổ sung signed streaming và telemetry chi tiết hơn.",
    ),
    (
        "AI có thể chấm sai thì sao?",
        "Hệ thống dùng rubric, schema JSON, normalizer, calibration, task relevance và lưu kết quả để audit. Điểm AI vẫn là hỗ trợ học tập, không thay chứng chỉ/giám khảo chính thức.",
    ),
    (
        "Speaking có thật sự chấm phát âm từ audio?",
        "Hiện audio được ghi/lưu, nhưng response nêu rõ TRANSCRIPT_ESTIMATE; phần lớn điểm dựa transcript và thời lượng. Roadmap là ASR có timestamp/phoneme hoặc dịch vụ acoustic scoring.",
    ),
    (
        "Nếu Ollama lỗi khi nộp test có câu AI?",
        "Server trả 502/503, không tạo TestAttempt và không trừ hạt đậu. Người học có thể nộp lại khi dịch vụ phục hồi.",
    ),
    (
        "Anti-cheat có ngăn được Alt+Tab không?",
        "Website không thể chặn tuyệt đối trên mọi hệ điều hành. Hệ thống phát hiện fullscreen/visibility/focus/heartbeat, ghi incident và có thể đánh trượt sau 3 violation đếm được.",
    ),
    (
        "Tại sao không đếm mọi cảnh báo anti-cheat là vi phạm?",
        "Camera, face detector và heartbeat có thể có false positive. Code phân biệt INFO/WARNING/VIOLATION; chỉ hành vi rời bài rõ ràng mới cộng vi phạm tự động.",
    ),
    (
        "Tại sao khóa có người mua không xóa vật lý?",
        "Để giữ lịch sử thanh toán, doanh thu, quyền lợi và audit. Hệ thống chuyển LOCKED; chỉ khóa chưa phát sinh quan hệ mới xóa an toàn.",
    ),
    (
        "Review có đáng tin không?",
        "Chỉ user đã đạt test COURSE mới được review, mỗi user cập nhật một review. Điều này giảm đánh giá ảo từ người chưa trải nghiệm.",
    ),
    (
        "Điểm yếu lớn nhất của database hiện tại?",
        "Feedback đang chứa cả review và marker tiến độ/chứng nhận. Khi scale nên tách CourseReview, LessonProgress, CourseCompletion để có ràng buộc và query rõ hơn.",
    ),
    (
        "File upload có an toàn chưa?",
        "Đã kiểm MIME, magic bytes, size và tên do server sinh. Chưa có antivirus/object storage; production cần S3/R2, signed URL, malware scan và lifecycle policy.",
    ),
    (
        "Hoàn tiền và rút tiền đã tự động chưa?",
        "Workflow, kiểm tra số dư, trạng thái và audit đã có; chuyển tiền thực tế vẫn ngoài hệ thống. Đây là ranh giới rõ giữa MVP và tích hợp ngân hàng production.",
    ),
    (
        "Dữ liệu nào chứng minh người học hoạt động?",
        "Enrollment chứng minh quyền; Feedback/VideoWatchProgress chứng minh lesson; TestAttempt chứng minh test; AiAssessment chứng minh AI; LearningActivity tạo timeline/streak.",
    ),
    (
        "Tại sao tài liệu nói Writing 2 hạt thay vì 3?",
        "Nguồn sự thật là lib/ai-points.ts hiện tại: WRITING_AI_COST = 2 và SPEAKING_AI_COST = 7. Tài liệu cũ trong repo chưa cập nhật điểm này.",
    ),
    (
        "Hướng mở rộng ưu tiên là gì?",
        "Object storage, Redis rate limit, tách bảng progress/review, hàng đợi email/AI, acoustic scoring cho Speaking, observability tập trung và bổ sung test tích hợp end-to-end.",
    ),
]


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 15, TEAL),
        ("Heading 1", 21, NAVY),
        ("Heading 2", 15, TEAL),
        ("Heading 3", 12, BLUE),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].paragraph_format.left_indent = Cm(0.55)
    styles["List Bullet"].paragraph_format.first_line_indent = Cm(-0.3)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    styles["List Bullet 2"].paragraph_format.left_indent = Cm(1.1)
    styles["List Number"].paragraph_format.left_indent = Cm(0.65)
    styles["List Number"].paragraph_format.first_line_indent = Cm(-0.35)

    code_style = styles.add_style("Code Path", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string(MID_GRAY)
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.paragraph_format.left_indent = Cm(0.4)
    code_style.paragraph_format.space_after = Pt(2)

    small_style = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    small_style.font.name = "Aptos"
    small_style.font.size = Pt(8.5)
    small_style.font.color.rgb = RGBColor.from_string(MID_GRAY)
    small_style.paragraph_format.space_after = Pt(3)

    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = "FINNCENTER  •  TÀI LIỆU BẢO VỆ ĐỒ ÁN"
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
        footer = section.footer
        add_page_number(footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FINNCENTER")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BẢO VỆ ĐỒ ÁN\n")
    run = title.add_run("CHỨC NĂNG VÀ CƠ CHẾ HOẠT ĐỘNG")
    run.font.size = Pt(24)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Nền tảng học ngoại ngữ tích hợp AI")

    doc.add_paragraph()
    add_callout(
        doc,
        "Phạm vi phân tích",
        "Tài liệu được xây dựng từ mã nguồn trong workspace hiện tại: 41 trang, 92 API route, 40 Prisma model và 17 enum. Nội dung ưu tiên code hiện hành khi có khác biệt với tài liệu cũ.",
        fill=LIGHT_TEAL,
        title_color=TEAL,
    )
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tài liệu thuyết trình • Kịch bản demo • Câu hỏi phản biện")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph("Ngày tạo: 30/07/2026", style="Small Note")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3 = doc.add_paragraph("Nguồn: mã nguồn dự án FinnCenter tại D:\\Work\\final-pj", style="Small Note")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    doc.add_heading("Cách sử dụng tài liệu", level=1)
    add_bullets(
        doc,
        [
            "Phần I–III dùng để mở đầu: vấn đề, kiến trúc, vai trò và dữ liệu.",
            "Phần IV là nội dung chính: bảo vệ từng nhóm chức năng bằng luồng, cơ chế, quy tắc và bằng chứng mã nguồn.",
            "Phần V là kịch bản demo 15–20 phút, ưu tiên một luồng end-to-end thay vì mở rời rạc mọi trang.",
            "Phần VI là bộ câu hỏi phản biện ngắn để luyện trả lời.",
            "Phần VII trình bày hạn chế và roadmap một cách trung thực.",
        ],
    )
    add_callout(
        doc,
        "Nguyên tắc khi bảo vệ",
        "Luôn trả lời theo cấu trúc: người dùng làm gì → API kiểm tra gì → dữ liệu nào thay đổi → cách chống lỗi/gian lận → giới hạn hiện tại.",
        fill=LIGHT_ORANGE,
        title_color=ORANGE,
    )

    doc.add_heading("Mục lục", level=1)
    add_matrix_table(
        doc,
        ["Phần", "Nội dung"],
        [
            ["I", "Tổng quan đồ án"],
            ["II", "Kiến trúc và cơ chế nền"],
            ["III", "Mô hình dữ liệu"],
            ["IV", "Bảo vệ từng chức năng và cơ chế"],
            ["V", "Kịch bản demo bảo vệ 15–20 phút"],
            ["VI", "Câu hỏi phản biện và trả lời ngắn"],
            ["VII", "Hạn chế và roadmap"],
            ["Phụ lục", "Bản đồ route, tệp nguồn ưu tiên và kết luận 60 giây"],
        ],
        widths=[1.15, 6.25],
    )
    doc.add_page_break()


def add_overview(doc: Document) -> None:
    doc.add_heading("PHẦN I. TỔNG QUAN ĐỒ ÁN", level=1)
    doc.add_heading("1. Bài toán và giá trị cốt lõi", level=2)
    add_paragraph(
        doc,
        "FinnCenter là nền tảng học ngoại ngữ tích hợp marketplace khóa học, LMS, bài kiểm tra, AI Writing/Speaking, thanh toán, quy trình giảng viên và quản trị. Hệ thống phục vụ ba vai trò STUDENT, TEACHER và ADMIN trên cùng một nguồn dữ liệu.",
    )
    add_matrix_table(
        doc,
        ["Vấn đề", "Cách FinnCenter giải quyết", "Bằng chứng đầu ra"],
        [
            ["Nội dung học rời rạc", "Course → Module → Lesson → Test theo learning gate", "Tiến độ, nextAction, lịch sử học"],
            ["Khó đánh giá kỹ năng", "Test chuẩn + AI Writing/Speaking theo rubric", "TestAttempt và AiAssessment"],
            ["Thiếu quy trình giảng viên", "Chứng chỉ + entrance test + anti-cheat + admin review", "Application log, attempt, anti-cheat log"],
            ["Thanh toán/doanh thu khó đối soát", "VNPay + ledger + revenue split + withdrawal workflow", "Payment, OrderItem, PointTransaction"],
            ["Admin thiếu dữ liệu quyết định", "Analytics, ranking, refund/report/security/email", "Dashboard và file export"],
        ],
        widths=[1.55, 3.65, 2.25],
    )
    doc.add_heading("2. Quy mô mã nguồn hiện tại", level=2)
    add_matrix_table(
        doc,
        ["Hạng mục", "Số lượng", "Ý nghĩa"],
        [
            ["Trang App Router", "41", "Khu public, auth, student, teacher, admin"],
            ["API route file", "92", "Route handlers cho nghiệp vụ và tích hợp"],
            ["Prisma model", "40", "Tài khoản, học tập, AI, tài chính, audit"],
            ["Enum", "17", "Role, status, kind, mode và state machine"],
        ],
        widths=[2.0, 1.0, 4.4],
    )
    doc.add_heading("3. Điểm khác biệt cần nói đúng theo code hiện tại", level=2)
    add_bullets(
        doc,
        [
            "Ví VND nội bộ không còn hoạt động; mua khóa và mua hạt đậu đều thanh toán trực tiếp qua VNPay.",
            "Writing AI feedback chi tiết hiện có giá 2 hạt; Speaking là 7 hạt.",
            "Học tập đã khóa tuần tự theo module: lesson → module test → module tiếp theo → final test.",
            "Bài test cho phép lưu mọi lượt làm và hiện cấu hình số lượt không giới hạn.",
            "Speaking lưu audio nhưng chấm hiện tại chủ yếu trên transcript (TRANSCRIPT_ESTIMATE).",
        ],
    )
    doc.add_page_break()


def add_architecture(doc: Document) -> None:
    doc.add_heading("PHẦN II. KIẾN TRÚC VÀ CƠ CHẾ NỀN", level=1)
    doc.add_heading("1. Kiến trúc tổng thể", level=2)
    add_matrix_table(
        doc,
        ["Lớp", "Công nghệ/thành phần", "Trách nhiệm"],
        [
            ["Trình duyệt", "React 19, Client Components, Web Worker, MediaRecorder", "Tương tác, form, ghi âm, timer, anti-cheat signals"],
            ["Web/App", "Next.js 16 App Router, Server Components", "Routing, SSR, điều phối trang và API"],
            ["Nghiệp vụ", "lib/*.ts và lib/ai/*.ts", "Auth, gate, payment, AI, revenue, validation"],
            ["Dữ liệu", "Prisma 7 + PostgreSQL", "Transaction, ràng buộc, lưu lịch sử/audit"],
            ["Tích hợp", "VNPay, Ollama, SMTP/Nodemailer", "Thanh toán, chấm AI, email"],
            ["Tệp", "public/uploads local", "Video, audio, ảnh, chứng chỉ, tài liệu"],
        ],
        widths=[1.0, 2.75, 3.7],
    )
    add_callout(
        doc,
        "Luồng chuẩn của một request ghi",
        "UI gửi dữ liệu → Route xác thực/validate → thư viện nghiệp vụ tính toán → Prisma transaction ghi dữ liệu → Notification/Email → UI đọc lại trạng thái.",
        fill=LIGHT_BLUE,
        title_color=BLUE,
    )

    doc.add_heading("2. Ma trận phân quyền", level=2)
    add_matrix_table(
        doc,
        ["Chức năng", "STUDENT", "TEACHER", "ADMIN"],
        [
            ["Xem marketplace", "Có", "Có", "Có"],
            ["Mua/học khóa", "Có", "Có; học thử khóa sở hữu", "Preview"],
            ["Tạo khóa/test", "Không", "Khóa sở hữu/ngôn ngữ duyệt", "Có"],
            ["Writing/Speaking AI", "Có; feedback trả hạt", "Có; feedback trả hạt", "Miễn phí"],
            ["Nộp hồ sơ giảng viên", "Có", "Có thể nộp lại theo workflow", "Duyệt"],
            ["Refund/report", "Tạo theo quyền", "Xử lý report khóa sở hữu", "Quyết định cuối"],
            ["Analytics/user admin", "Không", "Không", "Có"],
        ],
        widths=[2.15, 1.75, 2.1, 1.45],
    )

    doc.add_heading("3. Các cơ chế xuyên suốt", level=2)
    add_bullets(
        doc,
        [
            "Defense in depth: quyền được kiểm tra ở server dù giao diện đã ẩn nút.",
            "Idempotency: txnRef, sourceKey, unique constraints và upsert ngăn ghi trùng.",
            "State machine: Payment, Course, Application, Refund, Withdrawal, Report có trạng thái và chuyển trạng thái hợp lệ.",
            "Audit: lưu raw payment response, application log, anti-cheat log, bank change log và EmailLog.",
            "Transaction: nghiệp vụ tài chính/trừ điểm/hoàn tiền dùng transaction; các đoạn nhạy cảm dùng Serializable.",
            "Validation: dữ liệu đầu vào, file signature, thời gian, ownership và readiness đều được xác minh.",
        ],
    )
    doc.add_page_break()


def add_data_model(doc: Document) -> None:
    doc.add_heading("PHẦN III. MÔ HÌNH DỮ LIỆU", level=1)
    doc.add_heading("1. Nhóm dữ liệu chính", level=2)
    add_matrix_table(
        doc,
        ["Nhóm", "Model tiêu biểu", "Vai trò"],
        [
            ["Danh tính", "User, Session, EmailVerificationOtp, PasswordResetOtp, RegistrationSecurityEvent", "Tài khoản, xác thực, thu hồi và audit"],
            ["Học tập", "Course, Module, Lesson, VideoWatchProgress, Enrollment, Feedback, LearningActivity", "Nội dung, quyền và tiến độ"],
            ["Kiểm tra", "Test, Question, Answer, TestAttempt, CheatingLog", "Đề, câu hỏi, lượt thi và kết quả"],
            ["AI", "AiAssessment, PointTransaction, SystemSetting", "Kết quả AI, ledger hạt, cấu hình"],
            ["Giảng viên", "TeacherApplication, TeacherCertificate, TeacherEntranceQuestionInstance, AntiCheatLog, SuspiciousEvent", "Onboarding, thi tuần tự, giám sát"],
            ["Tài chính", "Payment, Order, OrderItem, CourseRefundRequest, TeacherBankAccount, TeacherRevenueWithdrawal, Complaint", "Thanh toán, chia/hoàn/rút tiền"],
            ["Hỗ trợ", "Notification, EmailLog, CourseReport, LearningLanguage", "Thông báo, báo cáo, danh mục"],
        ],
        widths=[1.2, 3.85, 2.35],
    )
    doc.add_heading("2. Quan hệ quan trọng khi trình bày", level=2)
    add_bullets(
        doc,
        [
            "User — Enrollment — Course: quyền học và mỗi user chỉ enroll một lần/khóa.",
            "Course — Module — Lesson và Course/Module — Test: cấu trúc learning gate.",
            "Test — Question — Answer — TestAttempt: snapshot lượt thi và chấm điểm.",
            "User — AiAssessment / PointTransaction / LearningActivity: kết quả AI, chi phí và lịch sử.",
            "TeacherApplication — QuestionInstance / AntiCheatLog / Certificate: bằng chứng xét duyệt.",
            "Payment — Order — OrderItem — RefundRequest: vòng đời tiền mua khóa.",
            "OrderItem.teacherRevenue — Withdrawal — Complaint: vòng đời doanh thu giảng viên.",
        ],
    )
    add_callout(
        doc,
        "Điểm kỹ thuật cần thừa nhận",
        "Feedback hiện đồng thời chứa review và marker PROGRESS/LESSON_START/COURSE_COMPLETED/CERT_SENT. Cách này nhanh cho MVP nhưng nên tách bảng chuyên biệt khi dữ liệu lớn.",
        fill=LIGHT_ORANGE,
        title_color=ORANGE,
    )
    doc.add_page_break()


def add_features(doc: Document) -> None:
    doc.add_heading("PHẦN IV. BẢO VỆ TỪNG CHỨC NĂNG VÀ CƠ CHẾ", level=1)
    add_paragraph(
        doc,
        "Mỗi mục dưới đây có thể trình bày trong 30–60 giây. Khi hội đồng hỏi sâu, dùng phần “Bằng chứng trong mã nguồn” để chỉ đúng route/thư viện.",
    )
    for index, feature in enumerate(FEATURES, start=1):
        add_feature_section(doc, index, feature)


def add_demo_script(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("PHẦN V. KỊCH BẢN DEMO BẢO VỆ 15–20 PHÚT", level=1)
    doc.add_heading("1. Luồng demo chính: từ học viên đến doanh thu", level=2)
    add_numbered(
        doc,
        [
            "Mở trang chủ/marketplace, tìm và lọc một khóa ACTIVE; giải thích public read và course status.",
            "Đăng nhập học viên; mở chi tiết khóa và bắt đầu mua. Nếu không muốn giao dịch thật, dùng khóa miễn phí hoặc dữ liệu Payment đã có để giải thích VNPay callback.",
            "Vào Khóa học của tôi; chỉ ra enrollment ACTIVE, tiến độ và nextAction.",
            "Mở một lesson văn bản/video; giải thích start/heartbeat/complete và điều kiện server.",
            "Hoàn tất module; cho thấy module test được mở, attempt token và chấm điểm.",
            "Mở Writing AI: chạy score-only trước, sau đó giải thích detailed feedback trừ 2 hạt trong transaction.",
            "Mở Speaking AI: ghi âm ngắn; nói rõ transcript estimate và audio lưu để xem lại.",
            "Chuyển sang teacher: tạo/chỉnh khóa, module, lesson/test; chỉ ra readiness 100 điểm và trạng thái pending.",
            "Mở doanh thu: giải thích 70/30, refund exclusion, bank OTP và withdrawal Serializable.",
            "Chuyển admin: duyệt khóa/refund/withdrawal và mở analytics để kết thúc.",
        ],
    )
    doc.add_heading("2. Luồng demo phụ: đăng ký giảng viên và anti-cheat", level=2)
    add_numbered(
        doc,
        [
            "Mở teacher registration, chọn ngôn ngữ và chứng chỉ.",
            "Giải thích camera → fullscreen → consent → server startedAt.",
            "Reveal một câu Speaking/Writing; chỉ thời gian chuẩn bị/trả lời và transition token.",
            "Thoát fullscreen một lần; chỉ log/violationCount và incident de-duplication.",
            "Mở màn admin xem application, điểm và anti-cheat log.",
        ],
    )
    doc.add_heading("3. Phương án khi demo lỗi", level=2)
    add_matrix_table(
        doc,
        ["Rủi ro", "Cách xử lý khi trình bày", "Thông điệp kỹ thuật"],
        [
            ["Ollama chậm/lỗi", "Dùng kết quả AiAssessment có sẵn", "Không lưu/trừ điểm khi AI lỗi"],
            ["VNPay không dùng sandbox", "Dùng Payment/OrderItem có sẵn và code callback", "Chữ ký + amount + idempotency"],
            ["Microphone/camera bị chặn", "Dùng audio/log đã lưu", "Feature phụ thuộc browser permission"],
            ["Email SMTP lỗi", "Mở EmailLog FAILED/SENT", "Email không rollback nghiệp vụ"],
            ["Dữ liệu demo thiếu", "Dùng seed/demo data hoặc record đã có", "Tránh sửa DB trực tiếp khi đang bảo vệ"],
        ],
        widths=[1.55, 2.65, 3.25],
    )
    add_callout(
        doc,
        "Mẹo trình bày",
        "Không cố demo 41 trang. Hãy demo một chuỗi dữ liệu xuyên vai trò và dùng tài liệu này để trả lời các chức năng còn lại.",
        fill=LIGHT_TEAL,
        title_color=TEAL,
    )


def add_qa(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("PHẦN VI. CÂU HỎI PHẢN BIỆN VÀ TRẢ LỜI NGẮN", level=1)
    for index, (question, answer) in enumerate(QA_ITEMS, start=1):
        p = doc.add_paragraph()
        set_keep_with_next(p)
        r = p.add_run(f"{index}. {question}")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.35)
        p2.add_run(answer)


def add_limitations(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("PHẦN VII. HẠN CHẾ VÀ ROADMAP", level=1)
    add_matrix_table(
        doc,
        ["Hiện trạng", "Rủi ro", "Hướng nâng cấp"],
        [
            ["Upload lưu public local", "Mất/không đồng bộ khi nhiều instance", "S3/R2, signed URL, antivirus, CDN"],
            ["Rate limit dùng Map trong RAM", "Không chia sẻ giữa instance", "Redis/Upstash, trusted proxy"],
            ["Speaking chấm transcript estimate", "Chưa đo phát âm âm vị trực tiếp", "ASR timestamp + acoustic/phoneme scoring"],
            ["AI phụ thuộc Ollama", "Timeout/quá tải", "Queue, circuit breaker, fallback model, monitoring"],
            ["Feedback overloaded", "Query/ràng buộc khó khi lớn", "Tách CourseReview, LessonProgress, Completion"],
            ["Chuyển/hoàn tiền thủ công", "Cần đối soát con người", "Bank payout/refund API + reconciliation"],
            ["Một số UI cấu hình chưa hoàn chỉnh", "Admin phải gọi API", "Hoàn thiện form cấu hình/language"],
            ["Một số chuỗi mojibake", "Trải nghiệm và bảo trì", "Chuẩn hóa toàn bộ repo UTF-8"],
            ["Thiếu E2E toàn luồng", "Regression nghiệp vụ", "Playwright + integration test DB/VNPay mock"],
        ],
        widths=[2.05, 2.35, 3.05],
    )
    doc.add_heading("Ưu tiên triển khai", level=2)
    add_numbered(
        doc,
        [
            "Object storage và Redis rate limit để sẵn sàng nhiều instance.",
            "Tách bảng progress/review và bổ sung migration dữ liệu.",
            "Hàng đợi cho AI/email, retry có kiểm soát và observability tập trung.",
            "Nâng Speaking từ transcript estimate lên phân tích âm thanh.",
            "Tự động hóa payout/refund và đối soát thanh toán.",
            "Bổ sung test tích hợp cho auth, learning gate, payment, refund, point ledger và withdrawal.",
        ],
    )
    add_callout(
        doc,
        "Cách nói khi bị hỏi về hạn chế",
        "Không né hạn chế. Nêu rõ: hệ thống hiện giải quyết đến đâu, rủi ro còn lại là gì, và thiết kế hiện tại đã chừa điểm nối nào cho bước nâng cấp.",
        fill=LIGHT_ORANGE,
        title_color=ORANGE,
    )


def add_appendix(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("PHỤ LỤC A. BẢN ĐỒ ROUTE TRÌNH BÀY NHANH", level=1)
    add_matrix_table(
        doc,
        ["Nhóm", "Trang chính", "API chính"],
        [
            ["Public/Auth", "/, /courses, /teachers, /auth/*", "/api/auth/*, /api/courses"],
            ["Student LMS", "/student, /my-courses, /student/hoc-bai", "/api/learning/*, /api/student/tests/*"],
            ["AI", "/student/writing-ai, /student/speaking-ai, /student/rewards", "/api/ai/*"],
            ["Teacher", "/teacher, /teacher/courses, /teacher/tests, /teacher/revenue", "/api/teacher/*"],
            ["Onboarding", "/teacher-registration", "/api/teacher-applications/*"],
            ["Admin", "/admin", "/api/admin/*"],
            ["Support", "/profile, /top-students, /about", "/api/profile, /api/notifications, /api/health"],
        ],
        widths=[1.2, 3.05, 3.25],
    )
    doc.add_heading("PHỤ LỤC B. TỆP NGUỒN ƯU TIÊN KHI HỘI ĐỒNG HỎI SÂU", level=1)
    for path in [
        "prisma/schema.prisma — mô hình dữ liệu và ràng buộc",
        "lib/auth.ts — token, cookie, mật khẩu và RBAC",
        "lib/course-learning-gates.ts — thứ tự học/module/test",
        "app/api/learning/lessons/[lessonId]/complete/route.ts — xác minh hoàn thành lesson",
        "lib/course-payment.ts và lib/vnpay.ts — thanh toán khóa",
        "lib/ai-points.ts — ledger hạt đậu và mua điểm",
        "app/api/student/tests/[testId]/submit/route.ts — chấm/lưu bài test",
        "app/api/ai/essay-evaluation/route.ts — Writing AI",
        "app/api/ai/speaking-evaluation/route.ts — Speaking AI",
        "lib/teacher-anti-cheat.ts và lib/teacher-sequential-exam.ts — thi đầu vào",
        "lib/teacher-revenue.ts — doanh thu khả dụng",
        "lib/admin-analytics.ts — KPI và ranking",
    ]:
        p = doc.add_paragraph(style="Code Path")
        p.add_run(path)

    doc.add_heading("PHỤ LỤC C. KẾT LUẬN 60 GIÂY", level=1)
    add_callout(
        doc,
        "Bài nói kết luận",
        "FinnCenter không chỉ là website hiển thị khóa học. Đồ án xây một vòng đời hoàn chỉnh: xác thực người dùng, mua khóa, học tuần tự có kiểm chứng, làm test và AI, hoàn thành/chứng nhận, đánh giá, onboarding giảng viên, quản lý nội dung, doanh thu và quản trị. Các nghiệp vụ quan trọng đều có kiểm tra phía server, trạng thái, transaction và audit. Hạn chế hiện tại tập trung ở hạ tầng production như object storage, distributed rate limit và phân tích âm thanh Speaking; các điểm này đã được xác định rõ trong roadmap.",
        fill=LIGHT_TEAL,
        title_color=TEAL,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— HẾT —")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(NAVY)


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Bảo vệ đồ án FinnCenter — Chức năng và cơ chế hoạt động"
    props.subject = "Tài liệu thuyết trình, kịch bản demo và câu hỏi phản biện"
    props.author = "Codex — tổng hợp từ mã nguồn dự án FinnCenter"
    props.keywords = "FinnCenter, đồ án, Next.js, Prisma, AI, VNPay, LMS"
    props.comments = "Nguồn sự thật: mã nguồn workspace tại thời điểm tạo tài liệu."


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    set_core_properties(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_overview(doc)
    add_architecture(doc)
    add_data_model(doc)
    add_features(doc)
    add_demo_script(doc)
    add_qa(doc)
    add_limitations(doc)
    add_appendix(doc)

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            set_keep_with_next(paragraph)
        if paragraph.style.name in {"Code Path", "Small Note"}:
            set_keep_together(paragraph)

    doc.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
